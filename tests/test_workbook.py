"""Workbook parsing: the real NHS file, and synthetic files with missing data."""

import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from navigator.workbook import (
    WorkbookError,
    classify_cell,
    default_organisation,
    parse_workbook,
)

ROOT = Path(__file__).resolve().parent.parent
REAL = sorted(ROOT.glob("Community-health-services-waiting-lists-*.xlsx"))

GSTT_BANDS = [38, 30, 76, 394, 175, 195, 1, 0]


# ------------------------------------------------------------- real workbook

@pytest.fixture(scope="module")
def real():
    if not REAL:
        pytest.skip("NHS workbook not present")
    return parse_workbook(REAL[-1].read_bytes())


def test_real_workbook_structure(real):
    assert real.ok, [c.message for c in real.checks]
    assert real.total.name == "Table 4"
    assert real.total.header_row == 4  # row 5 in Excel
    assert real.total.service_header == "(CYP) Community paediatric service"
    assert [s.name for s in real.bands] == [f"Table 4{x}" for x in "abcdefgh"]
    assert [b.band.lower_weeks for b in real.bands] == [0, 1, 2, 4, 12, 18, 52, 104]
    assert real.bands[-1].band.upper_weeks is None
    assert real.period == "August 2026"
    assert dict(real.cover)["Published"] == "08/10/2026"


def test_real_default_is_gstt(real):
    org = default_organisation(real.selectable_organisations())
    assert org.code == "RJ1"
    assert "GUY'S AND ST THOMAS'" in org.name


def test_real_gstt_figures(real):
    org = default_organisation(real.selectable_organisations())
    figures = real.figures_for(org)
    assert figures.total.value == 909
    assert [cell.value for _, cell, _ in figures.bands] == GSTT_BANDS
    assert figures.band_sum() == 909
    assert figures.waiting_over(52) == 1
    assert figures.waiting_over(104) == 0


def test_real_england_and_order(real):
    orgs = real.selectable_organisations()
    assert orgs[0].label == "England (national)"
    assert real.figures_for(orgs[0]).total.value == 174537
    assert all(o.org_type == "Provider" for o in orgs[1:])


def test_real_blank_provider_not_zero(real):
    blanks = [
        o for o in real.selectable_organisations()
        if real.figures_for(o).total.kind == "blank"
    ]
    assert blanks, "expected some providers with no paediatric submission"
    figures = real.figures_for(blanks[0])
    assert figures.total.value is None
    assert figures.waiting_over(52) is None
    assert figures.band_sum() is None


def test_real_notes(real):
    assert real.blank_note == "Blanks indicate no organisation submitted for the service line."
    assert "*" in real.marker_notes
    assert "sum of weekly breakdowns" in real.sum_note


# ------------------------------------------------------------- synthetic workbooks

BANDS = [
    "0-1 weeks (0-7 days)",
    ">1-2 weeks (8-14 days)",
    ">2-4 weeks (15-28 days)",
    ">4-12 weeks (29-84 days)",
    ">12-18 weeks (85-126 days)",
    ">18-52 weeks (127-364 days)",
    ">52-104 weeks (365-728 days)",
    "over 104 weeks (over 2 years)",
]


def make_workbook(values, header_offset=0, shuffle=False, include_paed=True, notes=None):
    """values: {org name: [total, b1..b8]}; None means a blank cell."""
    wb = Workbook()
    wb.remove(wb.active)
    cover = wb.create_sheet("Cover page")
    cover.append(["Title:", "Synthetic test workbook"])
    cover.append(["Source:", "Test"])
    titles = ["Total waiting list by service, March 2027"] + [
        f"Number waiting {b} by service, March 2027" for b in BANDS
    ]
    for index, title in enumerate(titles):
        ws = wb.create_sheet(f"Sheet{index}")
        ws.append([title])
        for note in notes or ["Blanks indicate no organisation submitted for the service line."]:
            ws.append([note])
        for _ in range(header_offset):
            ws.append([])
        headers = ["Organisation Type", "Organisation Code", "Organisation Name", "(A) Audiology"]
        if include_paed:
            headers.append("(CYP) Community paediatric service")
        if shuffle:
            headers = [headers[2], headers[-1], headers[0], headers[3], headers[1]]
        ws.append(headers)
        for name, row in values.items():
            record = {
                "Organisation Type": "National" if name == "England" else "Provider",
                "Organisation Code": "" if name == "England" else name[:3].upper(),
                "Organisation Name": name,
                "(A) Audiology": 5,
                "(CYP) Community paediatric service": row[index],
            }
            ws.append([record.get(h) for h in headers])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_blank_marker_and_zero_are_distinct():
    data = make_workbook(
        {
            "England": [10, 1, 1, 1, 1, 1, 1, 2, 2],
            "Blank Trust": [None] * 9,
            "Marker Trust": ["*", 1, 1, 1, 1, 1, 1, "*", "-"],
            "Zero Trust": [0, 0, 0, 0, 0, 0, 0, 0, 0],
            "Text Trust": ["1,234", "N/A", "", " ", "0", 1, 1, 1, 1],
        },
        notes=[
            "Blanks indicate no organisation submitted for the service line.",
            "* Suppressed to protect confidentiality.",
        ],
    )
    parsed = parse_workbook(data)
    assert parsed.ok, [c.message for c in parsed.checks]
    by_name = {o.name: o for o in parsed.organisations()}

    blank = parsed.figures_for(by_name["Blank Trust"])
    assert blank.total.kind == "blank" and blank.total.value is None
    assert all(cell.kind == "blank" for _, cell, _ in blank.bands)
    assert blank.waiting_over(52) is None

    marker = parsed.figures_for(by_name["Marker Trust"])
    assert marker.total.kind == "marker" and marker.total.raw == "*"
    assert marker.bands[7][1].raw == "-"
    assert marker.waiting_over(52) is None
    assert marker.waiting_over(104) is None
    assert parsed.marker_notes == {"*": "* Suppressed to protect confidentiality."}
    assert "-" not in parsed.marker_notes  # undocumented markers get no invented meaning

    zero = parsed.figures_for(by_name["Zero Trust"])
    assert zero.total.kind == "number" and zero.total.value == 0
    assert zero.waiting_over(52) == 0

    text = parsed.figures_for(by_name["Text Trust"])
    assert text.total.value == 1234
    assert text.bands[0][1].kind == "marker" and text.bands[0][1].raw == "N/A"
    assert text.bands[1][1].kind == "blank"
    assert text.bands[2][1].kind == "blank"
    assert text.bands[3][1].kind == "number" and text.bands[3][1].value == 0


def test_header_found_when_layout_moves():
    data = make_workbook({"England": [3, 0, 0, 0, 0, 0, 0, 1, 2]}, header_offset=4, shuffle=True)
    parsed = parse_workbook(data)
    assert parsed.ok
    assert parsed.total.header_row == 6
    england = parsed.figures_for(parsed.selectable_organisations()[0])
    assert england.total.value == 3
    assert england.waiting_over(52) == 3
    assert parsed.period == "March 2027"


def test_missing_paediatric_column_is_an_error():
    data = make_workbook({"England": [1] * 9}, include_paed=False)
    parsed = parse_workbook(data)
    assert not parsed.ok
    assert any("community paediatric" in c.message for c in parsed.checks if c.status == "error")


def test_band_sum_mismatch_is_reported_not_corrected():
    data = make_workbook({"England": [100, 1, 1, 1, 1, 1, 1, 1, 1]})
    parsed = parse_workbook(data)
    figures = parsed.figures_for(parsed.selectable_organisations()[0])
    assert figures.total.value == 100
    assert figures.band_sum() == 8


def test_not_an_excel_file():
    with pytest.raises(WorkbookError):
        parse_workbook(b"this is not a spreadsheet")


@pytest.mark.parametrize(
    "raw, kind, value",
    [(None, "blank", None), ("", "blank", None), (float("nan"), "blank", None),
     (0, "number", 0), ("12", "number", 12), ("1,500", "number", 1500),
     ("*", "marker", None), ("-", "marker", None), ("c", "marker", None)],
)
def test_classify_cell(raw, kind, value):
    cell = classify_cell(raw)
    assert cell.kind == kind
    assert cell.value == value
